import os
import re
import sys
import tkinter as tk
from tkinter import filedialog
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from colorama import init, Fore, Style

init(autoreset=True)

# ----------------- إعدادات الوورد ----------------- #
def set_bilingual_font(run, eng_font='Times New Roman', ar_font='Arial'):
    run.font.name = eng_font
    r = run._r
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), eng_font)
    rFonts.set(qn('w:hAnsi'), eng_font)
    rFonts.set(qn('w:cs'), ar_font)

def add_paragraph_border(paragraph, color="000000"):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    for border_name in ['w:top', 'w:left', 'w:bottom', 'w:right']:
        border = OxmlElement(border_name)
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '12')
        border.set(qn('w:space'), '6')
        border.set(qn('w:color'), color)
    pbdr.append(border)
    pPr.append(pbdr)

def add_page_number(doc):
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = "PAGE"
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)

# ----------------- معالجة المستند ----------------- #
def process_document(input_file, output_file, features):
    doc = docx.Document(input_file)
    point_counter = 1
    
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text: continue

        # 5. ترقيم النقاط
        if features['5'] and (text.startswith('-') or text.startswith('*')):
            p.text = f"{point_counter}- " + text[1:].strip()
            point_counter += 1
            
        # 1. تضليل ملاحظات
        if features['1'] and "ملاحظة" in text:
            p.text = "💡 " + p.text if not p.text.startswith("💡") else p.text
            for run in p.runs:
                run.font.bold = True
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                
        # 2. تمييز العناوين
        elif features['2'] and ("مهم" in text or "قاعدة" in text or p.style.name.startswith('Heading')):
            color = "CC0000" if "مهم" in text else "003366"
            add_paragraph_border(p, color=color)
            for run in p.runs:
                run.font.bold = True

        # 3. و 6. تخصيص الخطوط وتلوين الإنجليزي
        for run in p.runs:
            if features['3']:
                set_bilingual_font(run)
            
            if features['6'] and re.search(r'[A-Za-z]', run.text):
                # تلوين الكلمات الإنجليزية باللون الأزرق
                run.font.color.rgb = RGBColor(0, 0, 255)

    if features['4']:
        add_page_number(doc)

    doc.save(output_file)

# ----------------- الواجهة التفاعلية ----------------- #
def clear_screen():
    os.system('cls' if os.name == 'nt' else 'clear')

def display_banner():
    print(Fore.CYAN + "="*55)
    print(Fore.CYAN + "         مــكــتــبــة الــمــصــطــفــى")
    print(Fore.CYAN + "="*55)

def main():
    root = tk.Tk()
    root.withdraw()

    # الميزات: افتراضياً غير مفعلة
    features = {
        '1': {'name': 'تضليل الملاحظات', 'active': False},
        '2': {'name': 'تمييز العناوين', 'active': False},
        '3': {'name': 'الخط (عربي: Arial, إنجليزي: Times New Roman)', 'active': False},
        '4': {'name': 'ترقيم الصفحات', 'active': False},
        '5': {'name': 'ترقيم النقاط التلقائي', 'active': False},
        '6': {'name': 'تلوين الكلمات الإنجليزية', 'active': False}
    }

    while True:
        clear_screen()
        display_banner()
        print("\nقم بإدخال رقم الميزة لتفعيلها/إلغائها، أو اضغط 0 لاختيار الملف والبدء:\n")
        
        for key, value in features.items():
            color = Fore.GREEN if value['active'] else Fore.RED
            status = "[مفعل]" if value['active'] else "[معطل]"
            print(f"{color}{key}. {value['name']} - {status}")
            
        print(Fore.WHITE + "\n0. بدء المعالجة (اختيار الملف)")
        print(Fore.WHITE + "Q. خروج")
        
        choice = input("\n> ")
        
        if choice.lower() == 'q':
            sys.exit()
        elif choice == '0':
            break
        elif choice in features:
            features[choice]['active'] = not features[choice]['active']

    print(Fore.YELLOW + "\nيرجى اختيار ملف الوورد من النافذة المنبثقة...")
    input_file = filedialog.askopenfilename(
        title="اختر ملف الوورد الأصلي",
        filetypes=[("Word Documents", "*.docx")]
    )
    
    if not input_file:
        print(Fore.RED + "لم يتم اختيار أي ملف. تم الإلغاء.")
        return

    file_dir, file_name = os.path.split(input_file)
    name_only, extension = os.path.splitext(file_name)
    output_file = os.path.join(file_dir, f"{name_only}_معدل{extension}")
    
    print(Fore.CYAN + "\nجاري التعديل...")
    
    active_flags = {k: v['active'] for k, v in features.items()}
    try:
        process_document(input_file, output_file, active_flags)
        print(Fore.GREEN + "\nتم الانتهاء بنجاح!")
        print(Fore.GREEN + f"مسار الملف الجديد:\n{output_file}")
    except Exception as e:
        print(Fore.RED + f"\nحدث خطأ: {e}")
        
    input("\nاضغط Enter للخروج...")

if __name__ == "__main__":
    main()
